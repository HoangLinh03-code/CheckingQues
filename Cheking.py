"""
CheckDe.py - Chương trình check đề thi tự động với AI (IMPROVED VERSION)
- Kiểm tra CẤU TRÚC đầy đủ
- Kiểm tra NỘI DUNG chính xác
- Linh hoạt với nhiều định dạng
"""

import sys
import os
import re
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PyQt5.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QLabel, QFileDialog, QMessageBox, QTextEdit, QProgressBar,
    QGroupBox, QLineEdit, QRadioButton, QButtonGroup
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont, QColor
from dotenv import load_dotenv
from google.oauth2 import service_account
import traceback
from typing import Dict, List, Tuple
import time

sys.path.append(os.path.dirname(__file__))
from api.callApi import VertexClient

load_dotenv()

# ==================== LOAD CREDENTIALS ====================
def get_credentials():
    """Load Google Cloud credentials từ .env"""
    try:
        if getattr(sys, 'frozen', False):
            base_path = sys._MEIPASS
        else:
            base_path = os.path.dirname(__file__)
        
        dotenv_path = os.path.join(base_path, '.env')
        load_dotenv(dotenv_path)
        
        service_account_data = {
            "type": os.getenv("TYPE"),
            "project_id": os.getenv("PROJECT_ID"),
            "private_key_id": os.getenv("PRIVATE_KEY_ID"),
            "private_key": os.getenv("PRIVATE_KEY").replace('\\n', '\n'),
            "client_email": os.getenv("CLIENT_EMAIL"),
            "client_id": os.getenv("CLIENT_ID", ""),
            "auth_uri": os.getenv("AUTH_URI"),
            "token_uri": os.getenv("TOKEN_URI"),
            "auth_provider_x509_cert_url": os.getenv("AUTH_PROVIDER_X509_CERT_URL"),
            "client_x509_cert_url": os.getenv("CLIENT_X509_CERT_URL"),
            "universe_domain": os.getenv("UNIVERSE_DOMAIN")
        }
        
        credentials = service_account.Credentials.from_service_account_info(
            service_account_data,
            scopes=["https://www.googleapis.com/auth/cloud-platform"]
        )
        project_id = os.getenv('PROJECT_ID')
        
        return credentials, project_id
    except Exception as e:
        raise Exception(f"Không thể load credentials: {str(e)}")


# ==================== DOCX PARSER (IMPROVED) ====================
class DocxParser:
    """Parse DOCX để lấy nội dung câu hỏi - Linh hoạt hơn"""
    
    @staticmethod
    def extract_text_from_docx(docx_path: str) -> str:
        """Trích xuất toàn bộ text từ DOCX"""
        try:
            doc = Document(docx_path)
            full_text = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    full_text.append(text)
            
            return "\n".join(full_text)
        except Exception as e:
            raise Exception(f"Lỗi đọc file {docx_path}: {str(e)}")
    
    @staticmethod
    def parse_questions(text: str) -> Dict[int, str]:
        """
        Parse text thành từng câu hỏi riêng biệt - LINH HOẠT HƠN
        Hỗ trợ nhiều format: "Câu 1:", "**Câu 1:**", "Câu 1.", etc.
        """
        questions = {}
        lines = text.split("\n")
        current_question_num = None
        current_question_lines = []
        
        # Regex linh hoạt hơn để bắt nhiều dạng tiêu đề
        pattern = r'^\s*\*?\*?\s*(?:\[.*?\])?\s*Câu\s+(\d+)\s*[:\.]?\s*\*?\*?'
        
        for line in lines:
            # Phát hiện tiêu đề câu hỏi (linh hoạt)
            match = re.match(pattern, line, re.IGNORECASE)
            if match:
                # Lưu câu hỏi cũ
                if current_question_num is not None and current_question_lines:
                    questions[current_question_num] = "\n".join(current_question_lines)
                
                # Bắt đầu câu hỏi mới
                current_question_num = int(match.group(1))
                current_question_lines = [line]
            else:
                # Tiếp tục câu hỏi hiện tại
                if current_question_num is not None:
                    current_question_lines.append(line)
        
        # Lưu câu hỏi cuối
        if current_question_num is not None and current_question_lines:
            questions[current_question_num] = "\n".join(current_question_lines)
        
        return questions


# ==================== AI CHECKER (IMPROVED) ====================
class AIQuestionChecker:
    """Sử dụng AI để check câu hỏi - CẢI TIẾN"""
    
    def __init__(self, project_id: str, creds, model_name: str = "gemini-2.5-pro"):
        self.client = VertexClient(project_id, creds, model_name)
    
    def check_question(self, question_text: str, prompt_content: str, 
                      pdf_content: str = None) -> Dict:
        """
        Check 1 câu hỏi với AI - CẢI TIẾN
        
        Returns:
            {
                'is_correct': bool,
                'structure_check': str,  # MỚI: kiểm tra cấu trúc
                'content_check': str,     # MỚI: kiểm tra nội dung
                'evaluation': str,
                'suggestions': str,
                'knowledge': str
            }
        """
        # Tạo prompt đầy đủ
        pdf_section = f"## TÀI LIỆU THAM KHẢO (PDF)\n{pdf_content}" if pdf_content else ""

        full_prompt = (
            f"{prompt_content}\n\n"
            "---\n\n"
            "## NỘI DUNG CẦN ĐÁNH GIÁ\n\n"
            f"{question_text}\n\n"
            f"{pdf_section}\n\n"
            "---\n\n"
            "## YÊU CẦU\n\n"
            "Hãy đánh giá câu hỏi trên theo QUY TRÌNH 2 BƯỚC:\n\n"
            "**BƯỚC 1**: Kiểm tra CẤU TRÚC (có đủ các thành phần không?)\n"
            "**BƯỚC 2**: Kiểm tra NỘI DUNG (kiến thức, đáp án, logic có đúng không?)\n\n"
            "Trả về theo FORMAT SAU (BẮT BUỘC):\n\n"
            "```\n"
            "IS_CORRECT: YES/NO\n"
            "####\n"
            "STRUCTURE_CHECK:\n"
            "[Kiểm tra từng thành phần: tiêu đề, nội dung, đáp án, lời giải, giải thích, kết luận]\n"
            "####\n"
            "CONTENT_CHECK:\n"
            "[Kiểm tra kiến thức, đáp án, logic - nếu sai thì chỉ rõ sai gì]\n"
            "####\n"
            "EVALUATION:\n"
            "[Đánh giá tổng hợp - nếu NO thì phải có 'Gợi ý sửa: [chi tiết]']\n"
            "####\n"
            "SUGGESTIONS:\n"
            "[Gợi ý làm bài cho học sinh - 2-4 dòng, ngắn gọn]\n"
            "####\n"
            "KNOWLEDGE:\n"
            "[Kiến thức liên quan - 3-5 dòng, gạch đầu dòng]\n"
            "```\n\n"
            "**LƯU Ý**:\n"
            "- Nếu THIẾU bất kỳ thành phần nào → IS_CORRECT: NO\n"
            "- Nếu SAI kiến thức/đáp án → IS_CORRECT: NO\n"
            "- Gợi ý sửa phải CỤ THỂ: sửa gì, ở đâu, thành gì\n"
            "- NGẮN GỌN, SÚC TÍCH, ĐÚNG TRỌNG TÂM\n"
        )
        
        try:
            response = self.client.send_data_to_check(
                prompt=full_prompt,
                temperature=0.3
            )
            
            # Parse response với cấu trúc mới
            result = self._parse_response(response)
            return result
            
        except Exception as e:
            return {
                'is_correct': False,
                'structure_check': 'Không thể kiểm tra',
                'content_check': f"Lỗi: {str(e)}",
                'evaluation': f"Lỗi khi check: {str(e)}",
                'suggestions': "Không thể phân tích",
                'knowledge': "Không có"
            }
    
    def _parse_response(self, response: str) -> Dict:
        """Parse response từ AI - CẢI TIẾN"""
        result = {
            'is_correct': False,
            'structure_check': '',
            'content_check': '',
            'evaluation': '',
            'suggestions': '',
            'knowledge': ''
        }
        
        # Parse IS_CORRECT
        if 'IS_CORRECT: YES' in response or 'IS_CORRECT:YES' in response:
            result['is_correct'] = True
        
        # Parse các phần khác
        parts = response.split('####')
        
        for part in parts:
            part = part.strip()
            
            if 'STRUCTURE_CHECK:' in part:
                result['structure_check'] = part.replace('STRUCTURE_CHECK:', '').strip()
            elif 'CONTENT_CHECK:' in part:
                result['content_check'] = part.replace('CONTENT_CHECK:', '').strip()
            elif 'EVALUATION:' in part:
                result['evaluation'] = part.replace('EVALUATION:', '').strip()
            elif 'SUGGESTIONS:' in part:
                result['suggestions'] = part.replace('SUGGESTIONS:', '').strip()
            elif 'KNOWLEDGE:' in part:
                result['knowledge'] = part.replace('KNOWLEDGE:', '').strip()
        
        return result


# ==================== DOCX WRITER (IMPROVED) ====================
class DocxWriter:
    """Ghi kết quả đánh giá vào DOCX - CẢI TIẾN"""
    
    @staticmethod
    def add_evaluation_table(doc: Document, question_num: int, 
                            evaluation_result: Dict):
        """Thêm bảng đánh giá 2x3 vào document - CẢI TIẾN"""
        
        # Thêm khoảng cách trước bảng
        doc.add_paragraph()
        
        # Tạo bảng 2x3
        table = doc.add_table(rows=2, cols=3)
        table.style = 'Table Grid'
        
        # Set border
        for row in table.rows:
            for cell in row.cells:
                DocxWriter._set_cell_border(cell)
        
        # Header row
        headers = ['Đánh giá câu hỏi', 'Gợi ý làm bài', 'Kiến thức liên quan']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            
            # Format header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(11)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Background color
            DocxWriter._set_cell_background(cell, "D3D3D3")
        
        # Content row - CẢI TIẾN: Hiển thị cả structure và content check
        evaluation_text = f"[CẤU TRÚC]\n{evaluation_result['structure_check']}\n\n[NỘI DUNG]\n{evaluation_result['content_check']}\n\n[TỔNG HỢP]\n{evaluation_result['evaluation']}"
        
        contents = [
            evaluation_text,
            evaluation_result['suggestions'],
            evaluation_result['knowledge']
        ]
        
        for i, content in enumerate(contents):
            cell = table.rows[1].cells[i]
            cell.text = content
            
            # Format content
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
        
        # Thêm khoảng cách sau bảng
        doc.add_paragraph()
    
    @staticmethod
    def _set_cell_border(cell, **kwargs):
        """Set border cho cell"""
        tc = cell._element
        tcPr = tc.get_or_add_tcPr()
        
        tcBorders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            edge_element = OxmlElement(f'w:{edge}')
            edge_element.set(qn('w:val'), 'single')
            edge_element.set(qn('w:sz'), '4')
            edge_element.set(qn('w:space'), '0')
            edge_element.set(qn('w:color'), '000000')
            tcBorders.append(edge_element)
        
        tcPr.append(tcBorders)
    
    @staticmethod
    def _set_cell_background(cell, color_hex):
        """Set màu nền cho cell"""
        cell_properties = cell._element.get_or_add_tcPr()
        cell_shading = OxmlElement('w:shd')
        cell_shading.set(qn('w:fill'), color_hex)
        cell_properties.append(cell_shading)
    
    @staticmethod
    def create_output_docx(input_docx: str, questions_results: Dict[int, Dict],
                          output_path: str):
        """
        Tạo file DOCX output với bảng đánh giá
        
        Args:
            input_docx: Đường dẫn file DOCX gốc
            questions_results: {question_num: evaluation_result}
            output_path: Đường dẫn file output
        """
        # Load document gốc
        doc = Document(input_docx)
        
        # Tạo document mới
        new_doc = Document()
        
        # Copy các paragraph và thêm bảng đánh giá
        current_question_num = None
        pattern = r'^\s*\*?\*?\s*(?:\[.*?\])?\s*Câu\s+(\d+)\s*[:\.]?\s*\*?\*?'
        
        for para in doc.paragraphs:
            text = para.text.strip()
            
            # Phát hiện câu hỏi mới
            match = re.match(pattern, text, re.IGNORECASE)
            
            if match:
                # Xử lý câu hỏi cũ
                if current_question_num and current_question_num in questions_results:
                    # Thêm bảng đánh giá cho câu cũ
                    DocxWriter.add_evaluation_table(
                        new_doc,
                        current_question_num,
                        questions_results[current_question_num]
                    )
                
                # Bắt đầu câu hỏi mới
                current_question_num = int(match.group(1))
            
            # Copy paragraph
            new_para = new_doc.add_paragraph(text)
            # Copy formatting
            if para.style:
                try:
                    new_para.style = para.style
                except:
                    pass
        
        # Xử lý câu hỏi cuối cùng
        if current_question_num and current_question_num in questions_results:
            DocxWriter.add_evaluation_table(
                new_doc,
                current_question_num,
                questions_results[current_question_num]
            )
        
        # Lưu file
        new_doc.save(output_path)


# ==================== PROCESSING THREAD ====================
class CheckThread(QThread):
    """Thread xử lý check đề"""
    
    progress = pyqtSignal(str)
    finished_signal = pyqtSignal(list)
    error_signal = pyqtSignal(str)
    
    def __init__(self, input_paths: List[str], prompt_path: str, 
                 pdf_path: str, project_id: str, creds):
        super().__init__()
        self.input_paths = input_paths
        self.prompt_path = prompt_path
        self.pdf_path = pdf_path
        self.project_id = project_id
        self.creds = creds
        self.stop_requested = False
    
    def stop(self):
        self.stop_requested = True
    
    def run(self):
        try:
            # Load prompt
            self.progress.emit("📄 Đang load prompt...")
            with open(self.prompt_path, 'r', encoding='utf-8') as f:
                prompt_content = f.read()
            
            # Load PDF (nếu có)
            pdf_content = None
            if self.pdf_path and os.path.exists(self.pdf_path):
                self.progress.emit("📑 Đang load PDF tham khảo...")
                # TODO: Implement PDF reader nếu cần
                pdf_content = None
            
            # Khởi tạo AI checker
            self.progress.emit("🤖 Đang khởi tạo AI...")
            checker = AIQuestionChecker(self.project_id, self.creds)
            
            # Tạo output folder
            output_folder = "output_check"
            os.makedirs(output_folder, exist_ok=True)
            
            output_files = []
            
            # Xử lý từng file
            for idx, docx_path in enumerate(self.input_paths):
                if self.stop_requested:
                    self.progress.emit("⏸️ Đã dừng theo yêu cầu")
                    break
                
                filename = os.path.basename(docx_path)
                self.progress.emit(f"\n{'='*60}")
                self.progress.emit(f"📝 [{idx+1}/{len(self.input_paths)}] {filename}")
                self.progress.emit(f"{'='*60}\n")
                
                try:
                    # Parse DOCX
                    self.progress.emit("📖 Đang đọc file...")
                    parser = DocxParser()
                    full_text = parser.extract_text_from_docx(docx_path)
                    questions = parser.parse_questions(full_text)
                    
                    self.progress.emit(f"✔️ Tìm thấy {len(questions)} câu hỏi\n")
                    
                    # Check từng câu
                    questions_results = {}
                    correct_count = 0
                    incorrect_count = 0
                    
                    for qnum in sorted(questions.keys()):
                        if self.stop_requested:
                            break
                        
                        self.progress.emit(f"🔍 Đang check Câu {qnum}...")
                        
                        result = checker.check_question(
                            questions[qnum],
                            prompt_content,
                            pdf_content
                        )
                        
                        questions_results[qnum] = result
                        
                        if result['is_correct']:
                            status = "✅ ĐÚNG"
                            correct_count += 1
                        else:
                            status = "❌ SAI"
                            incorrect_count += 1
                        
                        self.progress.emit(f"   ➜ {status}")
                        
                        # Hiển thị lý do nếu sai
                        if not result['is_correct']:
                            eval_short = result['evaluation'][:100] + "..." if len(result['evaluation']) > 100 else result['evaluation']
                            self.progress.emit(f"   📌 {eval_short}")
                        
                        # Delay để tránh rate limit
                        time.sleep(0.5)
                    
                    # Tạo file output
                    self.progress.emit("\n📊 Đang tạo file output...")
                    
                    output_filename = f"checked_{filename}"
                    output_path = os.path.join(output_folder, output_filename)
                    
                    DocxWriter.create_output_docx(
                        docx_path,
                        questions_results,
                        output_path
                    )
                    
                    output_files.append(output_path)
                    
                    self.progress.emit(f"✅ Hoàn thành: {output_filename}")
                    self.progress.emit(f"   📈 Thống kê: {correct_count} đúng / {incorrect_count} sai / {len(questions)} tổng\n")
                    
                except Exception as e:
                    self.progress.emit(f"❌ Lỗi xử lý {filename}: {str(e)}\n")
                    continue
            
            self.progress.emit("\n" + "="*60)
            self.progress.emit("🎉 HOÀN THÀNH TẤT CẢ!")
            self.progress.emit("="*60)
            self.progress.emit(f"\n📁 Folder output: {output_folder}")
            self.progress.emit(f"📊 Đã xử lý: {len(output_files)}/{len(self.input_paths)} file\n")
            
            self.finished_signal.emit(output_files)
            
        except Exception as e:
            self.error_signal.emit(f"Lỗi: {str(e)}\n{traceback.format_exc()}")


# ==================== MAIN WINDOW ====================
class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Check Đề Thi Tự Động - AI Powered (IMPROVED)")
        self.resize(1000, 700)
        
        # Variables
        self.input_mode = "file"
        self.input_paths = []
        self.prompt_path = ""
        self.pdf_path = ""
        self.check_thread = None
        
        # Load credentials
        try:
            self.creds, self.project_id = get_credentials()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", f"Không thể load credentials: {str(e)}")
            sys.exit(1)
        
        self.setup_ui()
        self.apply_styles()
    
    def setup_ui(self):
        main_layout = QHBoxLayout()
        
        # === LEFT PANEL ===
        left_panel = QVBoxLayout()
        
        # 1. Chọn chế độ input
        mode_group = QGroupBox("📂 Chế Độ Input")
        mode_layout = QVBoxLayout()
        
        self.radio_file = QRadioButton("Chọn 1 file DOCX")
        self.radio_folder = QRadioButton("Chọn folder chứa nhiều DOCX")
        self.radio_file.setChecked(True)
        
        self.radio_file.toggled.connect(lambda: self.set_input_mode("file"))
        self.radio_folder.toggled.connect(lambda: self.set_input_mode("folder"))
        
        mode_layout.addWidget(self.radio_file)
        mode_layout.addWidget(self.radio_folder)
        mode_group.setLayout(mode_layout)
        
        # 2. Chọn input
        input_group = QGroupBox("📑 Chọn Input")
        input_layout = QVBoxLayout()
        
        self.input_label = QLineEdit("Chưa chọn")
        self.input_label.setReadOnly(True)
        
        self.btn_select_input = QPushButton("Chọn File/Folder")
        self.btn_select_input.clicked.connect(self.select_input)
        
        input_layout.addWidget(QLabel("Input:"))
        input_layout.addWidget(self.input_label)
        input_layout.addWidget(self.btn_select_input)
        input_group.setLayout(input_layout)
        
        # 3. Chọn prompt
        prompt_group = QGroupBox("📋 Prompt Check")
        prompt_layout = QVBoxLayout()
        
        self.prompt_label = QLineEdit("Chưa chọn prompt")
        self.prompt_label.setReadOnly(True)
        
        self.btn_select_prompt = QPushButton("Chọn File Prompt (.txt)")
        self.btn_select_prompt.clicked.connect(self.select_prompt)
        
        self.btn_edit_prompt = QPushButton("Sửa Prompt")
        self.btn_edit_prompt.clicked.connect(self.edit_prompt)
        
        prompt_btn_layout = QHBoxLayout()
        prompt_btn_layout.addWidget(self.btn_select_prompt)
        prompt_btn_layout.addWidget(self.btn_edit_prompt)
        
        prompt_layout.addWidget(QLabel("Prompt:"))
        prompt_layout.addWidget(self.prompt_label)
        prompt_layout.addLayout(prompt_btn_layout)
        prompt_group.setLayout(prompt_layout)
        
        # 4. Chọn PDF (optional)
        pdf_group = QGroupBox("📚 PDF Tham Khảo (Tùy chọn)")
        pdf_layout = QVBoxLayout()
        
        self.pdf_label = QLineEdit("Không có")
        self.pdf_label.setReadOnly(True)
        
        self.btn_select_pdf = QPushButton("Chọn PDF")
        self.btn_select_pdf.clicked.connect(self.select_pdf)
        
        pdf_layout.addWidget(QLabel("PDF (Optional):"))
        pdf_layout.addWidget(self.pdf_label)
        pdf_layout.addWidget(self.btn_select_pdf)
        pdf_group.setLayout(pdf_layout)
        
        # 5. Control buttons
        control_layout = QHBoxLayout()
        
        self.btn_start = QPushButton("▶️ BẮT ĐẦU CHECK")
        self.btn_start.clicked.connect(self.start_check)
        self.btn_start.setStyleSheet("""
            QPushButton {
                background-color: #27ae60;
                color: white;
                font-size: 14pt;
                font-weight: bold;
                padding: 10px;
                border-radius: 5px;
            }
            QPushButton:hover {
                background-color: #229954;
            }
            QPushButton:disabled {
                background-color: #95a5a6;
            }
        """)
        
        self.btn_stop = QPushButton("⏸️ DỪNG")
        self.btn_stop.clicked.connect(self.stop_check)
        self.btn_stop.setEnabled(False)
        self.btn_stop.setStyleSheet("""
            QPushButton {
                background-color: #e74c3c;
                color: white;
                font-size: 12pt;
                font-weight: bold;
                padding: 10px;
                border-radius: 5px;
            }
            QPushButton:hover {
                background-color: #c0392b;
            }
        """)
        
        control_layout.addWidget(self.btn_start)
        control_layout.addWidget(self.btn_stop)
        
        # 6. Progress
        progress_group = QGroupBox("📊 Tiến Độ")
        progress_layout = QVBoxLayout()
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setMaximum(100)
        self.progress_bar.setValue(0)
        
        progress_layout.addWidget(self.progress_bar)
        progress_group.setLayout(progress_layout)
        
        # Add to left panel
        left_panel.addWidget(mode_group)
        left_panel.addWidget(input_group)
        left_panel.addWidget(prompt_group)
        left_panel.addWidget(pdf_group)
        left_panel.addLayout(control_layout)
        left_panel.addWidget(progress_group)
        left_panel.addStretch()
        
        # === RIGHT PANEL ===
        right_panel = QVBoxLayout()
        
        log_group = QGroupBox("📋 Log Xử Lý")
        log_layout = QVBoxLayout()
        
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setFont(QFont("Consolas", 9))
        
        log_layout.addWidget(self.log_text)
        log_group.setLayout(log_layout)
        
        right_panel.addWidget(log_group)
        
        # === MAIN LAYOUT ===
        left_widget = QWidget()
        left_widget.setLayout(left_panel)
        left_widget.setMaximumWidth(400)
        
        right_widget = QWidget()
        right_widget.setLayout(right_panel)
        
        main_layout.addWidget(left_widget)
        main_layout.addWidget(right_widget)
        
        self.setLayout(main_layout)
    
    def apply_styles(self):
        self.setStyleSheet("""
            QWidget {
                font-family: 'Segoe UI', Arial;
                font-size: 10pt;
            }
            QGroupBox {
                font-weight: bold;
                border: 2px solid #3498db;
                border-radius: 5px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px;
            }
            QPushButton {
                background-color: #3498db;
                color: white;
                border: none;
                padding: 8px 15px;
                border-radius: 4px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #2980b9;
            }
            QPushButton:pressed {
                background-color: #21618c;
            }
            QPushButton:disabled {
                background-color: #95a5a6;
            }
            QLineEdit, QTextEdit {
                border: 1px solid #bdc3c7;
                border-radius: 3px;
                padding: 5px;
            }
            QProgressBar {
                border: 1px solid #bdc3c7;
                border-radius: 5px;
                text-align: center;
                height: 25px;
            }
            QProgressBar::chunk {
                background-color: #2ecc71;
            }
        """)
    
    def set_input_mode(self, mode):
        self.input_mode = mode
        if mode == "file":
            self.btn_select_input.setText("Chọn File DOCX")
        else:
            self.btn_select_input.setText("Chọn Folder")
        self.input_label.setText("Chưa chọn")
        self.input_paths = []
    
    def select_input(self):
        if self.input_mode == "file":
            file_path, _ = QFileDialog.getOpenFileName(
                self, "Chọn file DOCX", "", "Word Files (*.docx)"
            )
            if file_path:
                self.input_paths = [file_path]
                self.input_label.setText(os.path.basename(file_path))
                self.log_text.append(f"✔️ Đã chọn file: {os.path.basename(file_path)}")
        else:
            folder_path = QFileDialog.getExistingDirectory(
                self, "Chọn folder chứa DOCX"
            )
            if folder_path:
                docx_files = glob.glob(os.path.join(folder_path, "*.docx"))
                if docx_files:
                    self.input_paths = docx_files
                    self.input_label.setText(f"{folder_path} ({len(docx_files)} files)")
                    self.log_text.append(f"✔️ Đã chọn folder: {len(docx_files)} file DOCX")
                else:
                    QMessageBox.warning(self, "Cảnh báo", "Không tìm thấy file DOCX nào trong folder!")
    
    def select_prompt(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Chọn file prompt", "", "Text Files (*.txt)"
        )
        if file_path:
            self.prompt_path = file_path
            self.prompt_label.setText(os.path.basename(file_path))
            self.log_text.append(f"✔️ Đã chọn prompt: {os.path.basename(file_path)}")
    
    def edit_prompt(self):
        if not self.prompt_path or not os.path.exists(self.prompt_path):
            QMessageBox.warning(self, "Cảnh báo", "Vui lòng chọn file prompt trước!")
            return
        
        # Mở file bằng notepad/text editor
        import platform
        if platform.system() == 'Windows':
            os.system(f'notepad "{self.prompt_path}"')
        elif platform.system() == 'Darwin':  # macOS
            os.system(f'open -e "{self.prompt_path}"')
        else:  # Linux
            os.system(f'xdg-open "{self.prompt_path}"')
        
        self.log_text.append(f"📝 Đã mở prompt để chỉnh sửa")
    
    def select_pdf(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Chọn file PDF tham khảo", "", "PDF Files (*.pdf)"
        )
        if file_path:
            self.pdf_path = file_path
            self.pdf_label.setText(os.path.basename(file_path))
            self.log_text.append(f"✔️ Đã chọn PDF: {os.path.basename(file_path)}")
    
    def start_check(self):
        # Validation
        if not self.input_paths:
            QMessageBox.warning(self, "Lỗi", "Vui lòng chọn file/folder input!")
            return
        
        if not self.prompt_path or not os.path.exists(self.prompt_path):
            QMessageBox.warning(self, "Lỗi", "Vui lòng chọn file prompt!")
            return
        
        # Disable UI
        self.btn_start.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.log_text.clear()
        self.progress_bar.setValue(0)
        
        # Start thread
        self.check_thread = CheckThread(
            self.input_paths,
            self.prompt_path,
            self.pdf_path,
            self.project_id,
            self.creds
        )
        
        self.check_thread.progress.connect(self.update_log)
        self.check_thread.finished_signal.connect(self.on_finished)
        self.check_thread.error_signal.connect(self.on_error)
        
        self.check_thread.start()
    
    def stop_check(self):
        if self.check_thread and self.check_thread.isRunning():
            reply = QMessageBox.question(
                self, "Xác nhận",
                "Bạn có chắc muốn dừng?",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.Yes:
                self.check_thread.stop()
                self.log_text.append("\n⏸️ Đang dừng...")
    
    def update_log(self, message):
        self.log_text.append(message)
        self.log_text.verticalScrollBar().setValue(
            self.log_text.verticalScrollBar().maximum()
        )
    
    def on_finished(self, output_files):
        self.btn_start.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.progress_bar.setValue(100)
        
        QMessageBox.information(
            self, "Hoàn thành",
            f"✔️ Đã xử lý xong {len(output_files)} file!\n\n"
            f"📁 Kết quả trong folder: output_check"
        )
    
    def on_error(self, error_msg):
        self.btn_start.setEnabled(True)
        self.btn_stop.setEnabled(False)
        
        QMessageBox.critical(self, "Lỗi", error_msg)


# ==================== MAIN ====================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    
    window = MainWindow()
    window.show()
    
    sys.exit(app.exec_())